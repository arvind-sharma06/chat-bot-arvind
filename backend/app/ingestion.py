from __future__ import annotations

import base64
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

from anthropic import Anthropic
from docx import Document


@dataclass
class Chunk:
    text: str
    source: str


class DocxIngestor:
    def __init__(self, api_key: str, model_name: str, use_image_understanding: bool = True) -> None:
        self.client = Anthropic(api_key=api_key) if api_key else None
        self.model_name = model_name
        self.use_image_understanding = use_image_understanding

    def parse_docx(self, path: Path) -> list[Chunk]:
        chunks: list[Chunk] = []
        doc = Document(path)

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            chunks.extend(self._chunk_text("\n".join(paragraphs), str(path)))

        table_rows: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    table_rows.append(" | ".join(cells))
        if table_rows:
            chunks.extend(self._chunk_text("\n".join(table_rows), str(path)))

        image_texts = self._extract_image_descriptions(path)
        for idx, text in enumerate(image_texts, start=1):
            chunks.extend(self._chunk_text(text, f"{path}#image-{idx}"))

        return chunks

    def _chunk_text(self, text: str, source: str, chunk_size: int = 1200, overlap: int = 150) -> list[Chunk]:
        if len(text) <= chunk_size:
            return [Chunk(text=text, source=source)]

        out: list[Chunk] = []
        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            out.append(Chunk(text=text[start:end], source=source))
            if end == len(text):
                break
            start = max(0, end - overlap)
        return out

    def _extract_image_descriptions(self, path: Path) -> list[str]:
        if not self.use_image_understanding:
            return []
        if self.client is None:
            return []

        descriptions: list[str] = []
        with zipfile.ZipFile(path, "r") as zf:
            media_files = [name for name in zf.namelist() if name.startswith("word/media/")]
            for media_name in media_files:
                image_bytes = zf.read(media_name)
                media_type = self._detect_media_type(media_name)
                if media_type is None:
                    continue
                description = self._describe_image(image_bytes, media_type)
                if description:
                    descriptions.append(f"Image extracted from {path.name}: {description}")
        return descriptions

    @staticmethod
    def _detect_media_type(name: str) -> str | None:
        lower = name.lower()
        if lower.endswith(".png"):
            return "image/png"
        if lower.endswith(".jpg") or lower.endswith(".jpeg"):
            return "image/jpeg"
        if lower.endswith(".webp"):
            return "image/webp"
        return None

    def _describe_image(self, image_bytes: bytes, media_type: str) -> str:
        try:
            encoded = base64.b64encode(image_bytes).decode("utf-8")
            message = self.client.messages.create(
                model=self.model_name,
                max_tokens=400,
                temperature=0,
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": "Extract all relevant text and key details from this image for enterprise onboarding/support documentation.",
                            },
                            {
                                "type": "image",
                                "source": {
                                    "type": "base64",
                                    "media_type": media_type,
                                    "data": encoded,
                                },
                            },
                        ],
                    }
                ],
            )
            blocks = message.content
            text_parts = [b.text for b in blocks if hasattr(b, "text")]
            return "\n".join(text_parts).strip()
        except Exception:
            return ""
